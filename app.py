"""Local statement review server. Run: uvicorn app:app --host 127.0.0.1."""
import csv
import hashlib
import hmac
import secrets
import time
from contextvars import ContextVar
import io
import json
import os
import re
import sqlite3
import threading
import tempfile
import msoffcrypto
import uuid
from cryptography.fernet import Fernet, InvalidToken
from concurrent.futures import ThreadPoolExecutor
from contextlib import asynccontextmanager, contextmanager
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pymupdf as fitz
import openpyxl
import xlrd
from fastapi import FastAPI, File, Form, HTTPException, UploadFile, Request
from fastapi.responses import FileResponse, StreamingResponse, Response, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

ROOT = Path(__file__).parent
DATA = Path(os.environ.get('LEDGER_DATA_DIR', ROOT / 'data'))
DATA.mkdir(parents=True, exist_ok=True)
UPLOADS = DATA / 'uploads'
UPLOADS.mkdir(exist_ok=True)
DB = DATA / 'ledger.sqlite3'
POOL = ThreadPoolExecutor(max_workers=2)
LOCK = threading.RLock()
CURRENT_USER = ContextVar("current_user", default=None)
SESSION_SECONDS = 30 * 24 * 60 * 60

def user_id():
    value = CURRENT_USER.get()
    if value is None:
        raise HTTPException(401, "Please sign in")
    return value

MAX_BYTES = int(os.environ.get('MAX_UPLOAD_MB', '250')) * 1024 * 1024

@contextmanager
def db():
    conn = sqlite3.connect(DB, timeout=60)
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA foreign_keys=ON')
    try:
        yield conn
        conn.commit()
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()

def password_cipher():
    key_path = DATA / '.password-key'
    with LOCK:
        if not key_path.exists():
            with key_path.open('xb') as key_file:
                os.chmod(key_path, 0o600)
                key_file.write(Fernet.generate_key())
        return Fernet(key_path.read_bytes())

def saved_password(doc):
    if not doc['password_token']:
        return ''
    try:
        return password_cipher().decrypt(doc['password_token'].encode()).decode()
    except InvalidToken:
        raise PasswordRequired('Saved password cannot be read. Enter the file password again.') from None

def init_db():
    with db() as c:
        c.execute('PRAGMA foreign_keys=OFF')
        c.execute('PRAGMA journal_mode=WAL')
        c.executescript('''
        CREATE TABLE IF NOT EXISTS terms(id INTEGER PRIMARY KEY, phrase TEXT NOT NULL COLLATE NOCASE UNIQUE, enabled INTEGER NOT NULL DEFAULT 1);
        CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY, name TEXT NOT NULL, path TEXT NOT NULL, size INTEGER, status TEXT, error TEXT, warning TEXT, created_at TEXT, pages INTEGER DEFAULT 0, rows INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS source_rows(id INTEGER PRIMARY KEY, document_id TEXT REFERENCES documents(id) ON DELETE CASCADE, page INTEGER, sheet TEXT, row_number INTEGER, location TEXT, description TEXT, amount REAL, currency TEXT, columns TEXT, bbox TEXT);
        CREATE INDEX IF NOT EXISTS rows_document ON source_rows(document_id);
        CREATE TABLE IF NOT EXISTS matches(id INTEGER PRIMARY KEY, source_id INTEGER REFERENCES source_rows(id) ON DELETE CASCADE, term_id INTEGER REFERENCES terms(id) ON DELETE CASCADE, reviewed INTEGER DEFAULT 0, UNIQUE(source_id,term_id));
        CREATE INDEX IF NOT EXISTS matches_source ON matches(source_id);
        ''')
        if 'password_token' not in {row['name'] for row in c.execute('PRAGMA table_info(documents)')}:
            c.execute("ALTER TABLE documents ADD COLUMN password_token TEXT")
        if not c.execute("SELECT 1 FROM sqlite_master WHERE name='initialized'").fetchone():
            c.executemany('INSERT OR IGNORE INTO terms(phrase) VALUES (?)', [('ACH DEBIT RETURN',), ('OVERDRAFT FEE',), ('SERVICE CHARGE',), ('RETURN ITEM FEE',)])
            c.execute('CREATE TABLE initialized(version INTEGER)')
        c.executescript("""
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, username TEXT NOT NULL COLLATE NOCASE UNIQUE, password_hash TEXT NOT NULL, created_at REAL NOT NULL);
        CREATE TABLE IF NOT EXISTS sessions(token_hash TEXT PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE, expires_at REAL NOT NULL);
        CREATE TABLE IF NOT EXISTS auth_attempts(ip TEXT NOT NULL, attempted_at REAL NOT NULL);
        """)
        if 'user_id' not in {r['name'] for r in c.execute('PRAGMA table_info(documents)')}:
            c.execute('ALTER TABLE documents ADD COLUMN user_id INTEGER REFERENCES users(id)')
        if 'user_id' not in {r['name'] for r in c.execute('PRAGMA table_info(terms)')}:
            c.executescript("""
            CREATE TABLE terms_new(id INTEGER PRIMARY KEY, phrase TEXT NOT NULL COLLATE NOCASE, enabled INTEGER NOT NULL DEFAULT 1, user_id INTEGER REFERENCES users(id), UNIQUE(user_id,phrase));
            INSERT INTO terms_new(id,phrase,enabled) SELECT id,phrase,enabled FROM terms;
            DROP TABLE terms;
            ALTER TABLE terms_new RENAME TO terms;
            """)
        c.execute('CREATE INDEX IF NOT EXISTS documents_owner ON documents(user_id)')
        pending = c.execute("SELECT id FROM documents WHERE status IN ('queued','processing')").fetchall()
        c.execute("UPDATE documents SET status='queued' WHERE status='processing'")
    for row in pending:
        POOL.submit(process_document, row['id'])

@asynccontextmanager
async def lifespan(app):
    init_db()
    yield

app = FastAPI(title='findData', lifespan=lifespan)

def hash_password(password, salt=None):
    salt = salt or secrets.token_hex(16)
    digest = hashlib.scrypt(password.encode(), salt=bytes.fromhex(salt), n=16384, r=8, p=1).hex()
    return salt + ':' + digest

class Credentials(BaseModel):
    username: str = Field(min_length=3, max_length=64)
    password: str = Field(min_length=8, max_length=256, repr=False)

@app.middleware('http')
async def authentication(request: Request, call_next):
    if not request.url.path.startswith('/api/'):
        return await call_next(request)
    if request.method not in ('GET','HEAD','OPTIONS'):
        origin = request.headers.get('origin')
        if request.headers.get('sec-fetch-site') == 'cross-site' or (origin and origin != str(request.base_url).rstrip('/')):
            return JSONResponse({'detail':'Cross-origin request denied'}, status_code=403)
    public = request.url.path in ('/api/auth/login','/api/auth/register','/api/auth/status')
    token = request.cookies.get('ledger_session', '')
    with db() as c:
        session = c.execute('SELECT user_id FROM sessions WHERE token_hash=? AND expires_at>?', (hashlib.sha256(token.encode()).hexdigest(),time.time())).fetchone() if token else None
    if not public and not session:
        return JSONResponse({'detail':'Please sign in'}, status_code=401, headers={'Cache-Control':'no-store'})
    context_token = CURRENT_USER.set(session['user_id'] if session else None)
    try:
        response = await call_next(request)
        response.headers['Cache-Control'] = 'no-store'
        return response
    finally:
        CURRENT_USER.reset(context_token)

def throttle(request):
    ip = request.client.host if request.client else 'unknown'
    with LOCK, db() as c:
        c.execute('DELETE FROM auth_attempts WHERE attempted_at<?', (time.time()-300,))
        count = c.execute('SELECT COUNT(*) FROM auth_attempts WHERE ip=?', (ip,)).fetchone()[0]
        if count >= 20:
            raise HTTPException(429, 'Too many attempts. Try again in five minutes.')
        c.execute('INSERT INTO auth_attempts VALUES (?,?)', (ip,time.time()))

def sign_in_response(id, username, request):
    token = secrets.token_urlsafe(32)
    with db() as c:
        c.execute('DELETE FROM sessions WHERE expires_at<=?', (time.time(),))
        c.execute('INSERT INTO sessions VALUES (?,?,?)', (hashlib.sha256(token.encode()).hexdigest(),id,time.time()+SESSION_SECONDS))
    response = JSONResponse({'username':username})
    response.set_cookie('ledger_session',token,max_age=SESSION_SECONDS,httponly=True,samesite='lax',secure=request.url.scheme=='https' or os.environ.get('COOKIE_SECURE')=='1',path='/')
    return response

@app.get('/api/auth/status')
def auth_status():
    with db() as c:
        return {'setup_required':not bool(c.execute('SELECT 1 FROM users LIMIT 1').fetchone())}

@app.post('/api/auth/register')
def register(body: Credentials, request: Request):
    throttle(request)
    username = body.username.strip()
    if not re.fullmatch(r'[A-Za-z0-9_.-]{3,64}',username):
        raise HTTPException(422,'Use 3–64 letters, numbers, dots, underscores or hyphens for your username.')
    digest = hash_password(body.password)
    try:
        with LOCK, db() as c:
            first = not c.execute('SELECT 1 FROM users LIMIT 1').fetchone()
            id = c.execute('INSERT INTO users(username,password_hash,created_at) VALUES (?,?,?)', (username,digest,time.time())).lastrowid
            if first:
                c.execute('UPDATE documents SET user_id=? WHERE user_id IS NULL', (id,))
                c.execute('UPDATE terms SET user_id=? WHERE user_id IS NULL', (id,))
            else:
                c.executemany('INSERT INTO terms(phrase,user_id) VALUES (?,?)', ((phrase,id) for phrase in ['ACH DEBIT RETURN','OVERDRAFT FEE','SERVICE CHARGE','RETURN ITEM FEE']))
    except sqlite3.IntegrityError:
        raise HTTPException(409, 'Username is already taken') from None
    return sign_in_response(id,username,request)

@app.post('/api/auth/login')
def login(body: Credentials, request: Request):
    throttle(request)
    with db() as c:
        account = c.execute('SELECT * FROM users WHERE username=? COLLATE NOCASE', (body.username.strip(),)).fetchone()
    stored = account['password_hash'] if account else '00'*16+':'+'00'*64
    computed = hash_password(body.password,stored.split(':')[0])
    if not hmac.compare_digest(computed,stored) or not account:
        raise HTTPException(401,'Incorrect username or password')
    return sign_in_response(account['id'],account['username'],request)

@app.get('/api/auth/me')
def me():
    with db() as c:
        return {'username':c.execute('SELECT username FROM users WHERE id=?',(user_id(),)).fetchone()['username']}

@app.post('/api/auth/logout')
def logout(request: Request):
    token = request.cookies.get('ledger_session','')
    with db() as c:
        c.execute('DELETE FROM sessions WHERE token_hash=?',(hashlib.sha256(token.encode()).hexdigest(),))
    response = JSONResponse({'ok':True})
    response.delete_cookie('ledger_session',path='/')
    return response

class TermCreate(BaseModel):
    phrase: str = Field(min_length=1, max_length=150)

class TermPatch(BaseModel):
    phrase: str | None = Field(default=None, min_length=1, max_length=150)
    enabled: bool | None = None

class RetryRequest(BaseModel):
    password: str = Field(default="", max_length=1024, repr=False)

class PasswordRequired(ValueError):
    pass

class BulkReview(BaseModel):
    ids: list[int] = Field(min_length=1)
    reviewed: bool = True

class ReviewPatch(BaseModel):
    reviewed: bool

def normalize(text):
    return re.sub(r'\s+', ' ', str(text)).strip().casefold()

def matches_phrase(text, phrase):
    return re.search(r'(?<!\w)' + re.escape(normalize(phrase)) + r'(?!\w)', normalize(text)) is not None

def money(value):
    s = str(value).strip()
    if not s or not re.search(r'\d', s):
        return None
    negative = s.startswith('(') or s.endswith('-') or bool(re.search(r'\bDR$', s, re.I))
    s = re.sub(r'(?i)\b(?:USD|INR|EUR|GBP|CAD|AUD|DR|CR)\b', '', s)
    s = re.sub(r'[$₹€£,\s()]', '', s).rstrip('-')
    try:
        value = Decimal(s)
        if not value.is_finite():
            return None
        return float(-abs(value) if negative else value)
    except InvalidOperation:
        return None

def financials(columns, description):
    amount = None
    for key, value in columns.items():
        k = normalize(key)
        if any(word in k for word in ('debit', 'withdrawal', 'money out', 'paid out')):
            candidate = money(value)
            if candidate is not None and candidate != 0:
                amount = -abs(candidate)
                break
    if amount is None:
        for key, value in columns.items():
            k = normalize(key)
            if k in ('amount', 'transaction amount', 'net amount', 'credit', 'deposit', 'credit amount'):
                amount = money(value)
                if amount is not None:
                    break
    currency = None
    for value in [str(v) for k, v in columns.items() if 'currency' in normalize(k)] + [description]:
        found = re.search(r'\b(USD|INR|EUR|GBP|CAD|AUD|JPY|CHF|SGD|AED)\b', value, re.I)
        if found:
            currency = found.group(1).upper()
            break
        for symbol, code in [('₹', 'INR'), ('€', 'EUR'), ('£', 'GBP')]:
            if symbol in value:
                currency = code
    return amount, currency

def sheet_rows(rows, sheet):
    headers = None
    for number, values in enumerate(rows, 1):
        values = ['' if v is None else str(v) for v in values]
        if not any(v.strip() for v in values):
            continue
        # Keep all leading statement metadata; recognize a header only when transaction labels are present.
        lowered = [normalize(v) for v in values]
        is_header = any(v in ('description','narration','particulars','transaction description','details') for v in lowered) and any(any(w in v for w in ('date','amount','debit','credit','withdrawal')) for v in lowered)
        if is_header:
            headers = values
            continue
        keys = []
        for i in range(len(values)):
            key = headers[i] if headers and i < len(headers) and headers[i].strip() else f'Column {i+1}'
            if key in keys:
                key += f' ({i+1})'
            keys.append(key)
        columns = dict(zip(keys, values))
        description = ' | '.join(v for v in values if v)
        amount, currency = financials(columns, description)
        yield dict(page=None, sheet=sheet, row_number=number, location=f'{sheet}!A{number}:{openpyxl.utils.get_column_letter(len(values))}{number}', description=description, amount=amount, currency=currency, columns=columns, bbox=None)

def extract(path, progress, password=""):
    if path.suffix.lower() in ('.xlsx', '.xls'):
        with path.open('rb') as source:
            office = msoffcrypto.OfficeFile(source)
            if office.is_encrypted():
                if not password:
                    raise PasswordRequired('This file is password protected. Enter its password to continue.')
                with tempfile.TemporaryDirectory(prefix='findData-') as folder:
                    decrypted = Path(folder) / path.name
                    try:
                        office.load_key(password=password)
                        with decrypted.open('wb') as output:
                            office.decrypt(output)
                    except msoffcrypto.exceptions.InvalidKeyError:
                        raise PasswordRequired('Incorrect password. Please try again.') from None
                    yield from extract_content(decrypted, progress)
                return
    yield from extract_content(path, progress, password)

def extract_content(path, progress, password=""):
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        with fitz.open(path) as pdf:
            if pdf.needs_pass and not pdf.authenticate(password):
                raise PasswordRequired('Incorrect password. Please try again.' if password else 'This file is password protected. Enter its password to continue.')
            progress['pages'] = len(pdf)
            empty = []
            for page_index, page in enumerate(pdf):
                words = page.get_text('words', sort=True)
                if not words:
                    empty.append(page_index + 1)
                    continue
                table_regions = []
                for table in page.find_tables().tables:
                    extracted = table.extract()
                    if not extracted:
                        continue
                    header = extracted[0]
                    if not any(normalize(v or '') in ('description', 'narration', 'particulars', 'details', 'transaction description') for v in header):
                        continue
                    table_regions.append(fitz.Rect(table.bbox))
                    for entry in sheet_rows(extracted, 'PDF table'):
                        index = entry['row_number'] - 1
                        entry.update(page=page_index+1, sheet=None, location=f'Page {page_index+1}, table row {index+1}', bbox=list(table.rows[index].bbox))
                        yield entry
                words = [w for w in words if not any(region.contains(fitz.Point((w[0]+w[2])/2, (w[1]+w[3])/2)) for region in table_regions)]
                lines = []
                for word in words:
                    center = (word[1] + word[3]) / 2
                    nearby = next((line for line in reversed(lines[-6:]) if abs(line[0] - center) < 3), None)
                    if nearby is None:
                        nearby = [center, []]
                        lines.append(nearby)
                    nearby[1].append(word)
                for number, (_, line) in enumerate(lines, 1):
                    line.sort(key=lambda w: w[0])
                    text = ' '.join(w[4] for w in line)
                    bbox = [min(w[0] for w in line), min(w[1] for w in line), max(w[2] for w in line), max(w[3] for w in line)]
                    # PDF layout does not reliably identify debit/balance columns: leave amount unknown.
                    _, currency = financials({}, text)
                    yield dict(page=page_index+1, sheet=None, row_number=number, location=f'Page {page_index+1}, line {number}', description=text, amount=None, currency=currency, columns={'Full extracted row': text}, bbox=bbox)
            if empty:
                progress['warning'] = f'{len(empty)} page(s) have no extractable text and need OCR. First pages: {", ".join(map(str,empty[:20]))}.'
            else:
                progress['warning'] = 'PDF rows are reconstructed from text positions. Verify wrapped descriptions and amounts against the source; Amounts are classified only for recognized table columns.'
    elif suffix == '.xlsx':
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book:
                yield from sheet_rows(sheet.iter_rows(values_only=True), sheet.title)
        finally:
            book.close()
    elif suffix == '.xls':
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            for sheet in book.sheets():
                yield from sheet_rows((sheet.row_values(i) for i in range(sheet.nrows)), sheet.name)
        finally:
            book.release_resources()
    else:
        with path.open(encoding='utf-8-sig', newline='') as f:
            sample = f.read(8192)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            yield from sheet_rows(csv.reader(f, dialect), 'CSV')

def rescan(c, document_id=None):
    owner = c.execute('SELECT user_id FROM documents WHERE id=?', (document_id,)).fetchone()['user_id'] if document_id else user_id()
    terms = c.execute('SELECT * FROM terms WHERE enabled=1 AND user_id=?', (owner,)).fetchall()
    c.execute('DELETE FROM matches WHERE term_id IN (SELECT id FROM terms WHERE enabled=0 AND user_id=?)', (owner,))
    rows = c.execute('SELECT r.id,r.description FROM source_rows r JOIN documents d ON d.id=r.document_id WHERE d.user_id=? AND (? IS NULL OR d.id=?)', (owner,document_id,document_id))
    for row in rows:
        for term in terms:
            if matches_phrase(row['description'], term['phrase']):
                c.execute('INSERT OR IGNORE INTO matches(source_id,term_id) VALUES (?,?)', (row['id'],term['id']))
            else:
                c.execute('DELETE FROM matches WHERE source_id=? AND term_id=?', (row['id'],term['id']))

def process_document(document_id, password=""):
    try:
        with LOCK, db() as c:
            doc = c.execute('SELECT * FROM documents WHERE id=?', (document_id,)).fetchone()
            if not doc:
                return
            c.execute("UPDATE documents SET status='processing', error=NULL WHERE id=?", (document_id,))
            c.execute('DELETE FROM source_rows WHERE document_id=?', (document_id,))
        password = password or saved_password(doc)
        progress = {'pages':0, 'warning':None}
        batch = []
        count = 0
        def flush():
            with LOCK, db() as c:
                if not c.execute('SELECT 1 FROM documents WHERE id=?', (document_id,)).fetchone():
                    raise ValueError('Document removed')
                c.executemany('INSERT INTO source_rows(document_id,page,sheet,row_number,location,description,amount,currency,columns,bbox) VALUES (?,?,?,?,?,?,?,?,?,?)', batch)
                c.execute('UPDATE documents SET rows=? WHERE id=?', (count,document_id))
            batch.clear()
        for row in extract(Path(doc['path']), progress, password):
            batch.append((document_id,row['page'],row['sheet'],row['row_number'],row['location'],row['description'],row['amount'],row['currency'],json.dumps(row['columns']),json.dumps(row['bbox'])))
            count += 1
            if len(batch) >= 500:
                flush()
        if batch:
            flush()
        with LOCK, db() as c:
            rescan(c, document_id)
            if password:
                token = password_cipher().encrypt(password.encode()).decode()
                c.execute('UPDATE documents SET password_token=? WHERE id=?', (token,document_id))
            c.execute("UPDATE documents SET status='ready',pages=?,rows=?,warning=? WHERE id=?", (progress['pages'],count,progress['warning'],document_id))
    except Exception as exc:
        with LOCK, db() as c:
            c.execute('DELETE FROM source_rows WHERE document_id=?', (document_id,))
            c.execute("UPDATE documents SET status=?,error=?,rows=0 WHERE id=?", ("password_required" if isinstance(exc, PasswordRequired) else "error", str(exc)[:500],document_id))

def get_doc(c, id):
    row = c.execute('SELECT * FROM documents WHERE id=? AND user_id=?', (id,user_id())).fetchone()
    if row is None:
        raise HTTPException(404, 'Document not found')
    return row

@app.get('/api/terms')
def terms():
    with db() as c:
        return [dict(r) for r in c.execute('SELECT id,phrase,enabled FROM terms WHERE user_id=? ORDER BY id', (user_id(),))]

@app.post('/api/terms')
def add_term(body: TermCreate):
    phrase = ' '.join(body.phrase.split())
    if not phrase:
        raise HTTPException(422, 'Enter a search phrase')
    try:
        with LOCK, db() as c:
            id = c.execute('INSERT INTO terms(phrase,user_id) VALUES (?,?)', (phrase,user_id())).lastrowid
            rescan(c)
            return dict(c.execute('SELECT * FROM terms WHERE id=?', (id,)).fetchone())
    except sqlite3.IntegrityError:
        raise HTTPException(409, 'This search phrase already exists')

@app.patch('/api/terms/{id}')
def edit_term(id: int, body: TermPatch):
    try:
        with LOCK, db() as c:
            existing = c.execute('SELECT * FROM terms WHERE id=? AND user_id=?', (id,user_id())).fetchone()
            if not existing:
                raise HTTPException(404, 'Search phrase not found')
            if body.phrase is not None:
                phrase = ' '.join(body.phrase.split())
                if not phrase:
                    raise HTTPException(422, 'Enter a search phrase')
                c.execute('UPDATE terms SET phrase=? WHERE id=?', (phrase,id))
                if normalize(phrase) != normalize(existing['phrase']):
                    c.execute('DELETE FROM matches WHERE term_id=?', (id,))
            if body.enabled is not None:
                c.execute('UPDATE terms SET enabled=? WHERE id=?', (body.enabled,id))
            rescan(c)
            return {'ok':True}
    except sqlite3.IntegrityError:
        raise HTTPException(409, 'This search phrase already exists')

@app.delete('/api/terms/{id}')
def delete_term(id: int):
    with LOCK, db() as c:
        c.execute('DELETE FROM terms WHERE id=? AND user_id=?', (id,user_id()))
    return {'ok':True}

@app.post('/api/documents')
async def upload(files: list[UploadFile] = File(...), password: str = Form(default="", max_length=1024)):
    if len(files) > 20:
        raise HTTPException(400, 'Upload up to 20 files at a time')
    staged = []
    try:
        for file in files:
            name = Path(file.filename or 'unnamed').name
            suffix = Path(name).suffix.lower()
            if suffix not in ('.pdf','.xlsx','.xls','.csv'):
                raise HTTPException(400, f'Unsupported file: {name}')
            id = uuid.uuid4().hex
            path = UPLOADS / (id + suffix)
            staged.append((id,name,path,0))
            size = 0
            with path.open('wb') as out:
                while chunk := await file.read(1024 * 1024):
                    size += len(chunk)
                    if size > MAX_BYTES:
                        raise HTTPException(413, f'{name} exceeds {MAX_BYTES//1024//1024} MB limit')
                    out.write(chunk)
            if size == 0:
                raise HTTPException(400, f'{name} is empty')
            staged[-1] = (id,name,path,size)
        with db() as c:
            for id,name,path,size in staged:
                c.execute('INSERT INTO documents(id,name,path,size,status,created_at,user_id) VALUES (?,?,?,?,?,?,?)', (id,name,str(path),size,'queued',datetime.now(timezone.utc).isoformat(),user_id()))
    except BaseException:
        for _,_,path,_ in staged:
            path.unlink(missing_ok=True)
        raise
    finally:
        for file in files:
            await file.close()
    for id,_,_,_ in staged:
        POOL.submit(process_document,id,password)
    return {'documents':[{'id':id,'name':name,'status':'queued'} for id,name,_,_ in staged]}

@app.delete('/api/documents/{id}')
def remove_document(id: str):
    with LOCK, db() as c:
        doc = get_doc(c,id)
        c.execute('DELETE FROM documents WHERE id=?', (id,))
        Path(doc['path']).unlink(missing_ok=True)
    return {'ok':True}

@app.post('/api/documents/{id}/retry')
def retry(id: str, body: RetryRequest | None = None):
    with LOCK, db() as c:
        doc = get_doc(c,id)
        if doc['status'] in ('queued','processing'):
            raise HTTPException(409, 'Already processing')
        c.execute("UPDATE documents SET status='queued',error=NULL WHERE id=?", (id,))
    POOL.submit(process_document,id,body.password if body else '')
    return {'ok':True}

@app.get('/api/documents/{id}/file')
def original(id: str):
    with db() as c:
        doc = get_doc(c,id)
    return FileResponse(doc['path'], filename=doc['name'], content_disposition_type='inline' if doc['name'].lower().endswith('.pdf') else 'attachment')

@app.get('/api/documents/{id}/preview')
def preview(id: str):
    with db() as c:
        doc = get_doc(c,id)
    if not doc['name'].lower().endswith('.pdf'):
        raise HTTPException(400, 'Preview is available for PDFs only')
    with fitz.open(doc['path']) as pdf:
        if pdf.needs_pass:
            try:
                password = saved_password(doc)
            except PasswordRequired as exc:
                raise HTTPException(409, str(exc)) from None
            if not password or not pdf.authenticate(password):
                raise HTTPException(409, 'Enter the file password from Statements first')
            return Response(pdf.tobytes(encryption=fitz.PDF_ENCRYPT_NONE), media_type='application/pdf', headers={'Cache-Control':'no-store'})
    return FileResponse(doc['path'], media_type='application/pdf', headers={'Cache-Control':'no-store'})

@app.delete('/api/documents/{id}/password')
def forget_password(id: str):
    with LOCK, db() as c:
        doc = get_doc(c,id)
        if doc['status'] in ('queued','processing'):
            raise HTTPException(409, 'Wait for processing to finish before removing the saved password')
        c.execute('UPDATE documents SET password_token=NULL WHERE id=?', (id,))
    return {'ok':True}

@app.post('/api/matches/review')
def bulk_review(body: BulkReview):
    ids = set(body.ids)
    with LOCK, db() as c:
        existing = {row['id'] for row in c.execute('SELECT m.id FROM matches m JOIN source_rows r ON r.id=m.source_id JOIN documents d ON d.id=r.document_id WHERE d.user_id=?', (user_id(),))}
        if not ids.issubset(existing):
            raise HTTPException(409, 'Some matches have changed. Refresh and select them again.')
        c.executemany('UPDATE matches SET reviewed=? WHERE id=?', ((body.reviewed,id) for id in ids))
    return {'updated':len(ids)}

@app.patch('/api/matches/{id}')
def review(id: int, body: ReviewPatch):
    with db() as c:
        if not c.execute('UPDATE matches SET reviewed=? WHERE id=? AND source_id IN (SELECT r.id FROM source_rows r JOIN documents d ON d.id=r.document_id WHERE d.user_id=?)', (body.reviewed,id,user_id())).rowcount:
            raise HTTPException(404, 'Match not found')
    return {'ok':True}

def all_matches(c):
    for row in c.execute('''SELECT m.id,m.source_id,m.reviewed,t.phrase AS term,r.*,d.name AS document_name FROM matches m JOIN source_rows r ON r.id=m.source_id JOIN terms t ON t.id=m.term_id JOIN documents d ON d.id=r.document_id WHERE d.status='ready' AND d.user_id=? ORDER BY d.created_at DESC,r.row_number,m.id''', (user_id(),)):
        result = dict(row)
        result['columns'] = json.loads(result['columns'])
        result['bbox'] = json.loads(result['bbox'])
        result['reviewed'] = bool(result['reviewed'])
        yield result

@app.get('/api/dashboard')
def dashboard():
    with db() as c:
        documents = [dict(r) for r in c.execute('SELECT d.id,d.name,d.size,d.status,d.error,d.warning,d.created_at,d.pages,d.rows,(d.password_token IS NOT NULL) AS has_password,COUNT(m.id) AS match_count FROM documents d LEFT JOIN source_rows r ON r.document_id=d.id LEFT JOIN matches m ON m.source_id=r.id WHERE d.user_id=? GROUP BY d.id ORDER BY d.created_at DESC', (user_id(),))]
        matches = list(all_matches(c))
        term_list = [dict(r) for r in c.execute('SELECT id,phrase,enabled FROM terms WHERE user_id=? ORDER BY id', (user_id(),))]
        currencies = {}
        seen = set()
        for row in matches:
            if row['amount'] is not None and row['amount'] < 0:
                # Distinct source row, not distinct keyword, for charge totals.
                key = row['source_id']
                if key in seen:
                    continue
                seen.add(key)
                code = row['currency']
                if not code:
                    continue
                currencies[code] = round(currencies.get(code,0) + abs(row['amount']),2)
        return {'documents':documents,'matches':matches,'terms':term_list,'currencies':currencies,'stats':{'documents':len(documents),'matches':len(matches),'debit_total':None,'unreviewed':sum(not m['reviewed'] for m in matches)}}

@app.get('/api/export.csv')
def export(document_id: str = '', term: str = '', reviewed: bool | None = None, q: str = '', debit_only: bool = False):
    def generate():
        out = io.StringIO()
        writer = csv.writer(out)
        writer.writerow(['Document','Term','Page','Sheet','Row','Location','Description','Amount','Currency','Reviewed','Full row','PDF bounding box'])
        yield '\ufeff' + out.getvalue()
        with db() as c:
            for m in all_matches(c):
                if document_id and m['document_id'] != document_id or term and m['term'] != term or reviewed is not None and m['reviewed'] != reviewed or q and normalize(q) not in normalize(m['description']+' '+m['document_name']+' '+m['term']) or debit_only and (m['amount'] is None or m['amount'] >= 0):
                    continue
                values = [m['document_name'],m['term'],m['page'],m['sheet'],m['row_number'],m['location'],m['description'],m['amount'],m['currency'],m['reviewed'],json.dumps(m['columns']),json.dumps(m['bbox'])]
                values = ["'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@','\t','\r')) else v for v in values]
                out.seek(0)
                out.truncate(0)
                writer.writerow(values)
                yield out.getvalue()
    return StreamingResponse(generate(), media_type='text/csv', headers={'Content-Disposition':'attachment; filename="findData-matches.csv"'})

@app.post('/api/demo')
def demo():
    id = uuid.uuid4().hex
    path = UPLOADS / (id+'.csv')
    path.write_text('Date,Description,Debit,Credit,Balance,Currency\n2026-10-01,ACH DEBIT RETURN - Example supplier,35.00,,4965.00,USD\n2026-10-02,Client invoice payment,,1200.00,6165.00,USD\n2026-10-03,SERVICE CHARGE - Monthly account,12.00,,6153.00,USD\n2026-10-04,OVERDRAFT FEE,25.00,,6128.00,USD\n', encoding='utf-8')
    with db() as c:
        c.execute('INSERT INTO documents(id,name,path,size,status,created_at,user_id) VALUES (?,?,?,?,?,?,?)', (id,'Sample statement.csv',str(path),path.stat().st_size,'queued',datetime.now(timezone.utc).isoformat(),user_id()))
    POOL.submit(process_document,id)
    return {'id':id}

app.mount('/static', StaticFiles(directory=ROOT/'static'), name='assets')
app.mount('/', StaticFiles(directory=ROOT/'static', html=True), name='static')
